from openpyxl import load_workbook


def assign_missing_entity_ids(excel_path):
    """
    Assign permanent VET_###### IDs to rows in the Guide sheet
    that do not already have an entity_id.
    """

    workbook = load_workbook(excel_path)
    sheet = workbook["Guide"]

    # Find the entity_id column.
    headers = {
        cell.value: cell.column
        for cell in sheet[1]
        if cell.value is not None
    }

    entity_id_column = headers["entity_id"]

    # Find the highest existing VET number.
    highest_id = 0

    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(
            row=row,
            column=entity_id_column
        ).value

        if isinstance(value, str) and value.startswith("VET_"):
            try:
                number = int(value.split("_")[1])
                highest_id = max(highest_id, number)
            except ValueError:
                pass

    next_id = highest_id + 1

    # Assign IDs only to rows that already contain an entity.
    for row in range(2, sheet.max_row + 1):

        entity_cell = sheet.cell(
            row=row,
            column=entity_id_column
        )

        diagnosis = sheet.cell(
            row=row,
            column=headers["Diagnosis"]
        ).value

        if diagnosis and not entity_cell.value:

            entity_cell.value = f"VET_{next_id:06d}"

            next_id += 1

    workbook.save(excel_path)


if __name__ == "__main__":
    assign_missing_entity_ids(
        r"data\Pathophysiology Guide.xlsx"
    )

    print("Entity IDs assigned successfully.")